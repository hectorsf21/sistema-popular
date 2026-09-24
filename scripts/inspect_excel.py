import openpyxl
import sys
import os
import json

def inspect_excel():
    excel_path = "./data/padron.xlsx"
    if not os.path.exists(excel_path):
        print(f"Archivo no encontrado: {excel_path}")
        return

    print("Cargando libro de Excel...")
    # read_only=True para archivos grandes (67MB) para no agotar la memoria RAM
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    sheet_names = wb.sheetnames
    print(f"Hojas encontradas: {sheet_names}")

    ws = wb[sheet_names[0]]
    
    # Obtener encabezados de la primera fila
    headers = []
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        headers = [str(cell or '').strip() for cell in row]
        break

    print(f"Encabezados encontrados ({len(headers)} columnas):")
    for i, h in enumerate(headers):
        print(f"  Columna {i+1}: '{h}'")

    # Muestra de las primeras 5 filas
    print("\n--- MUESTRA DE PRIMERAS 5 FILAS ---")
    row_count = 0
    for row in ws.iter_rows(min_row=2, max_row=6, values_only=True):
        row_count += 1
        print(f"Fila {row_count}: {row}")

if __name__ == "__main__":
    inspect_excel()
