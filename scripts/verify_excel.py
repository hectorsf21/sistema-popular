import sys
import os
import json
import re
import openpyxl
from datetime import datetime

def normalize_cedula(cedula_str):
    if not cedula_str:
        return ""
    cleaned = re.sub(r'[\.\s\-]', '', str(cedula_str).upper())
    if cleaned.isdigit():
        return f"V{cleaned}"
    return cleaned

def format_date(val):
    if not val:
        return ""
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    match = re.search(r'(\d{4})[/-](\d{1,2})[/-](\d{1,2})', s)
    if match:
        return f"{match.group(1)}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"
    return s

def verify_person(cedula_input, fecha_nac_input=None, excel_path=None):
    if not excel_path:
        excel_path = os.getenv("EXCEL_PADRON_PATH", "./data/padron.xlsx")

    if not os.path.exists(excel_path):
        return {
            "success": False,
            "message": f"El archivo del padrón no se encuentra en la ruta: {excel_path}"
        }

    norm_target_cedula = normalize_cedula(cedula_input)

    try:
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]

        # Obtener encabezados
        headers = []
        for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
            headers = [str(c or '').strip().upper() for c in row]
            break

        # Buscar índices de columnas
        col_idx = {
            "NAC": 0,          # Column1 (V / E)
            "CI": 1,           # C.I
            "APE1": 2,         # PRIMER APELLIDO
            "APE2": 3,         # SEGUNDO APELLIDO
            "NOM1": 4,         # PRIMER NOMBRE
            "NOM2": 5,         # SEGUNDO NOMBRE
            "FECHA_NAC": 7,    # FECHA DE NACIMIENTO
            "MUNICIPIO": 12,   # NOMBRE DEL MUNICIPIO
            "PARROQUIA": 14,   # NOMBRE DE PARROQUIA
            "CV": 16,          # NOMBRE DEL CV
            "COMUNA": 18       # CIRCUITO COMUNAL
        }

        for idx, h in enumerate(headers):
            if 'C.I' in h or 'CEDULA' in h: col_idx["CI"] = idx
            elif 'PRIMER APELLIDO' in h: col_idx["APE1"] = idx
            elif 'SEGUNDO APELLIDO' in h: col_idx["APE2"] = idx
            elif 'PRIMER NOMBRE' in h: col_idx["NOM1"] = idx
            elif 'SEGUNDO NOMBRE' in h: col_idx["NOM2"] = idx
            elif 'FECHA DE NACIMIENTO' in h or 'FECHA_NAC' in h: col_idx["FECHA_NAC"] = idx
            elif 'NOMBRE DEL MUNICIPIO' in h or 'MUNICIPIO' in h: col_idx["MUNICIPIO"] = idx
            elif 'NOMBRE DE PARROQUIA' in h or 'PARROQUIA' in h: col_idx["PARROQUIA"] = idx
            elif 'CIRCUITO COMUNAL' in h or 'COMUNA' in h: col_idx["COMUNA"] = idx
            elif 'NOMBRE DEL CV' in h: col_idx["CV"] = idx

        found_person = None

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[col_idx["CI"]] is None:
                continue

            nac = str(row[col_idx["NAC"]] or 'V').strip().upper() if col_idx["NAC"] < len(row) else 'V'
            ci_raw = str(row[col_idx["CI"]]).strip()
            
            row_cedula = normalize_cedula(f"{nac}{ci_raw}")
            row_cedula_num_only = normalize_cedula(ci_raw)

            if row_cedula == norm_target_cedula or row_cedula_num_only == norm_target_cedula:
                dob_formatted = format_date(row[col_idx["FECHA_NAC"]] if col_idx["FECHA_NAC"] < len(row) else None)
                
                # Validar fecha de nacimiento si fue suministrada
                if fecha_nac_input and fecha_nac_input != "null":
                    target_dob_clean = re.sub(r'[^\d]', '', str(fecha_nac_input))
                    row_dob_clean = re.sub(r'[^\d]', '', str(dob_formatted))
                    
                    if target_dob_clean and row_dob_clean and target_dob_clean != row_dob_clean:
                        return {
                            "success": False,
                            "message": f"La fecha de nacimiento no coincide con el registro oficial del padrón."
                        }

                nom1 = str(row[col_idx["NOM1"]] or '').strip() if col_idx["NOM1"] < len(row) else ""
                nom2 = str(row[col_idx["NOM2"]] or '').strip() if col_idx["NOM2"] < len(row) else ""
                ape1 = str(row[col_idx["APE1"]] or '').strip() if col_idx["APE1"] < len(row) else ""
                ape2 = str(row[col_idx["APE2"]] or '').strip() if col_idx["APE2"] < len(row) else ""

                full_name = f"{nom1} {nom2} {ape1} {ape2}".replace("  ", " ").strip()

                municipio = str(row[col_idx["MUNICIPIO"]] or '').strip() if col_idx["MUNICIPIO"] < len(row) else ""
                parroquia = str(row[col_idx["PARROQUIA"]] or '').strip() if col_idx["PARROQUIA"] < len(row) else ""
                comuna_raw = str(row[col_idx["COMUNA"]] or '').strip() if col_idx["COMUNA"] < len(row) else ""
                cv_raw = str(row[col_idx["CV"]] or '').strip() if col_idx["CV"] < len(row) else ""

                comunidad = comuna_raw if comuna_raw and comuna_raw.upper() not in ['#N/A', 'NONE', 'N/A'] else (cv_raw or parroquia or municipio)

                found_person = {
                    "cedula": f"{nac}-{ci_raw}",
                    "nombre": full_name,
                    "nombres": f"{nom1} {nom2}".strip(),
                    "apellidos": f"{ape1} {ape2}".strip(),
                    "fechaNacimiento": dob_formatted,
                    "comunidad": comunidad,
                    "municipio": municipio,
                    "parroquia": parroquia,
                    "circuitoComunal": comuna_raw if comuna_raw.upper() not in ['#N/A', 'NONE'] else None
                }
                break

        if found_person:
            return {
                "success": True,
                "message": "Usuario verificado exitosamente en el padrón electoral oficial.",
                "data": found_person
            }
        else:
            return {
                "success": False,
                "message": f"La cédula {cedula_input} no aparece en el padrón electoral oficial."
            }

    except Exception as e:
        return {
            "success": False,
            "message": f"Error al procesar el archivo Excel del padrón: {str(e)}"
        }

if __name__ == "__main__":
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    
    args = sys.argv[1:]
    if not args:
        print(json.dumps({
            "success": False,
            "message": "Uso: verify_excel.py <cedula> [fecha_nacimiento] [excel_path]"
        }))
        sys.exit(1)

    cedula = args[0]
    fecha_nac = args[1] if len(args) > 1 and args[1] != "null" and args[1] != "" else None
    path = args[2] if len(args) > 2 and args[2] != "null" and args[2] != "" else None

    result = verify_person(cedula, fecha_nac, path)
    print(json.dumps(result, ensure_ascii=False))
