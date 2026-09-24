import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def generate_sample_padron():
    # Asegurar directorio data
    os.makedirs("./data", exist_ok=True)
    file_path = "./data/padron.xlsx"
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PadronElectoral"
    
    # Encabezados
    headers = ["CEDULA", "FECHA_NACIMIENTO", "NOMBRES", "APELLIDOS", "COMUNIDAD", "MUNICIPIO", "PARROQUIA", "ESTADO"]
    ws.append(headers)
    
    # Estilo de encabezados
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Datos de demostración / muestra
    sample_data = [
        # Jefes y Patrullados de prueba
        ["V-12345678", "1990-05-15", "Carlos Alberto", "Rodríguez Pérez", "Comunidad Bolívar Centro", "Bolívar", "Catedral", "Activo"],
        ["V-87654321", "1988-11-20", "María Elena", "Gómez López", "Comunidad Bolívar Centro", "Bolívar", "Catedral", "Activo"],
        ["V-11223344", "1995-02-10", "José Luis", "Fernández Silva", "Sector El Carmen", "Bolívar", "Agua Salada", "Activo"],
        ["V-55667788", "1992-08-30", "Ana Karina", "Martínez Díaz", "Sector El Carmen", "Bolívar", "Agua Salada", "Activo"],
        ["V-99887766", "1985-04-12", "Pedro Antonio", "Sánchez Morales", "Comunidad La Sabanita", "Bolívar", "La Sabanita", "Activo"],
        ["V-22334455", "1998-12-05", "Laura Beatriz", "Torres Ruiz", "Comunidad La Sabanita", "Bolívar", "La Sabanita", "Activo"],
        ["V-33445566", "1991-07-22", "Roberto José", "Ramírez Vargas", "Comunidad Bolívar Centro", "Bolívar", "Catedral", "Activo"],
        ["V-44556677", "1987-09-18", "Carmen Rosa", "Hernández Gil", "Sector El Carmen", "Bolívar", "Agua Salada", "Activo"],
        ["V-66778899", "1993-01-25", "Gabriel Eduardo", "Castillo Blanco", "Comunidad La Sabanita", "Bolívar", "La Sabanita", "Activo"],
        ["V-77889900", "1996-06-30", "Daniela Sofía", "Mendoza Castro", "Comunidad Bolívar Centro", "Bolívar", "Catedral", "Activo"],
        ["V-88990011", "1989-03-14", "Francisco Javier", "Rojas Medina", "Sector El Carmen", "Bolívar", "Agua Salada", "Activo"],
        ["V-99001122", "1994-10-08", "Patricia Isabel", "Romero Suárez", "Comunidad La Sabanita", "Bolívar", "La Sabanita", "Activo"],
        ["V-10101010", "1997-12-12", "Miguel Ángel", "Navarro Paredes", "Comunidad Bolívar Centro", "Bolívar", "Catedral", "Activo"],
        ["V-20202020", "1991-04-04", "Sofía Alejandra", "Acosta Guerrero", "Sector El Carmen", "Bolívar", "Agua Salada", "Activo"],
        ["V-30303030", "1986-08-08", "Alejandro José", "Pineda Rivas", "Comunidad La Sabanita", "Bolívar", "La Sabanita", "Activo"],
    ]
    
    # Generar más registros de prueba dinámicamente hasta 30 personas
    for i in range(1, 16):
        cedula_num = 15000000 + i * 1234
        cedula = f"V-{cedula_num}"
        dob = f"19{80 + (i % 20):02d}-{(i % 12) + 1:02d}-{(i % 28) + 1:02d}"
        nombres = f"UsuarioPrueba {i}"
        apellidos = f"ApellidoDemo {i}"
        comunidad = f"Comunidad Sector { (i % 5) + 1 }"
        municipio = "Bolívar"
        parroquia = "Vista Hermosa"
        sample_data.append([cedula, dob, nombres, apellidos, comunidad, municipio, parroquia, "Activo"])

    for row in sample_data:
        ws.append(row)
        
    # Ajustar ancho de columnas
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    wb.save(file_path)
    print(f"[SUCCESS] Excel del Padron de prueba generado exitosamente en: {file_path}")

if __name__ == "__main__":
    generate_sample_padron()
