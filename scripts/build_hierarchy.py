import openpyxl
import os
import json
import re

def clean_name(name):
    if not name or str(name).strip().upper() in ['#N/A', 'NONE', 'N/A', 'NAN', 'NULL']:
        return None
    s = str(name).strip()
    # Limpiar prefijos comunes si se prefiere o mantener legible
    return s

def build_hierarchy():
    excel_path = "./data/padron.xlsx"
    if not os.path.exists(excel_path):
        print(f"Error: Archivo no encontrado en {excel_path}")
        return

    print("Procesando archivo padron.xlsx para extraer arbol de municipios, parroquias y comunas...")
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    # Diccionario jerárquico: { municipio: { parroquia: set(comunas) } }
    tree = {}
    
    total_filas = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        total_filas += 1
        if total_filas % 50000 == 0:
            print(f"  Procesadas {total_filas} filas...")

        municipio = clean_name(row[12]) # Columna 13: NOMBRE DEL MUNICIPIO
        parroquia = clean_name(row[14]) # Columna 15: NOMBRE DE PARROQUIA
        comuna = clean_name(row[18])    # Columna 19: CIRCUITO COMUNAL

        if not municipio:
            continue

        if municipio not in tree:
            tree[municipio] = {}

        if parroquia:
            if parroquia not in tree[municipio]:
                tree[municipio][parroquia] = set()

            if comuna:
                tree[municipio][parroquia].add(comuna)

    print(f"Procesamiento completo. Total filas leídas: {total_filas}")

    # Convertir a estructura JSON / TypeScript limpia y ordenada
    result_tree = {}
    total_municipios = len(tree)
    total_parroquias = 0
    total_comunas = 0

    for mun_name, parroquias_map in sorted(tree.items()):
        result_tree[mun_name] = {}
        for par_name, comunas_set in sorted(parroquias_map.items()):
            total_parroquias += 1
            sorted_comunas = sorted(list(comunas_set))
            total_comunas += len(sorted_comunas)
            result_tree[mun_name][par_name] = sorted_comunas

    # Guardar en src/data/ubicaciones.json y src/data/ubicaciones.ts
    os.makedirs("./src/data", exist_ok=True)
    
    json_path = "./src/data/ubicaciones.json"
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(result_tree, f, ensure_ascii=False, indent=2)

    ts_path = "./src/data/ubicaciones.ts"
    with open(ts_path, "w", encoding="utf-8") as f:
        f.write("// ARCHIVO GENERADO AUTOMATICAMENTE DESDE EL PADRON EXCEL\n")
        f.write("// Estructura jerarquica: Municipio -> Parroquia -> Comunas / Circuitos Comunales\n\n")
        f.write("export interface UbicacionesTree {\n")
        f.write("  [municipio: string]: {\n")
        f.write("    [parroquia: string]: string[];\n")
        f.write("  };\n")
        f.write("}\n\n")
        f.write("export const UBICACIONES_DATA: UbicacionesTree = ")
        f.write(json.dumps(result_tree, ensure_ascii=False, indent=2))
        f.write(";\n\n")
        f.write("export const MUNICIPIOS_LIST = Object.keys(UBICACIONES_DATA);\n")

    print(f"[SUCCESS] Archivo JSON creado en: {json_path}")
    print(f"[SUCCESS] Archivo TS creado en: {ts_path}")
    print(f"[INFO] Resumen de Datos Extraidos:")
    print(f"   - Total Municipios: {total_municipios}")
    print(f"   - Total Parroquias: {total_parroquias}")
    print(f"   - Total Comunas / Circuitos Comunales: {total_comunas}")

if __name__ == "__main__":
    build_hierarchy()
