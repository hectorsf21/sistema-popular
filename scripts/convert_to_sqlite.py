import sqlite3
import openpyxl
import os
import re

excel_path = "./data/padron.xlsx"
db_path = "./data/padron.db"

if not os.path.exists(excel_path):
    print(f"Error: No se encontró el archivo {excel_path}")
    exit(1)

print("Creando base de datos SQLite indexada desde padron.xlsx...")
if os.path.exists(db_path):
    os.remove(db_path)

conn = sqlite3.connect(db_path)
cursor = conn.cursor()

cursor.execute('''
CREATE TABLE padron (
    cedula_norm TEXT PRIMARY KEY,
    cedula_full TEXT,
    nombre TEXT,
    nombres TEXT,
    apellidos TEXT,
    fecha_nac TEXT,
    comunidad TEXT,
    municipio TEXT,
    parroquia TEXT
)
''')

wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]

count = 0
batch = []

def clean_val(v):
    return str(v).strip() if v is not None else ""

for row in ws.iter_rows(min_row=2, values_only=True):
    if not row or row[1] is None:
        continue
    
    count += 1
    nac = clean_val(row[0]).upper() or 'V'
    ci = clean_val(row[1])
    
    # Normalizar quitando puntos, guiones y espacios
    norm = re.sub(r'[\.\s\-]', '', f"{nac}{ci}").upper()
    
    ape1 = clean_val(row[2])
    ape2 = clean_val(row[3])
    nom1 = clean_val(row[4])
    nom2 = clean_val(row[5])
    
    # Fecha de nacimiento
    dob = clean_val(row[7])
    if len(dob) >= 10:
        dob = dob[:10]
        
    mun = clean_val(row[12])
    par = clean_val(row[14])
    comuna = clean_val(row[18])
    cv = clean_val(row[16])
    
    comunidad = comuna if comuna and comuna.upper() not in ['#N/A', 'NONE', 'N/A'] else (cv or par or mun)
    full_name = f"{nom1} {nom2} {ape1} {ape2}".replace("  ", " ").strip()

    batch.append((
        norm,
        f"{nac}-{ci}",
        full_name,
        f"{nom1} {nom2}".strip(),
        f"{ape1} {ape2}".strip(),
        dob,
        comunidad,
        mun,
        par
    ))

    # Guardar en lotes de 10.000 para que sea súper rápido
    if len(batch) >= 10000:
        cursor.executemany('INSERT OR IGNORE INTO padron VALUES (?,?,?,?,?,?,?,?,?)', batch)
        conn.commit()
        batch = []
        print(f"Procesadas {count} filas...")

if batch:
    cursor.executemany('INSERT OR IGNORE INTO padron VALUES (?,?,?,?,?,?,?,?,?)', batch)
    conn.commit()

# Crear índice para búsquedas instantáneas
cursor.execute('CREATE INDEX IF NOT EXISTS idx_cedula ON padron (cedula_norm)')
conn.commit()
conn.close()

print(f"¡Éxito! Se indexaron {count} registros en {db_path}.")